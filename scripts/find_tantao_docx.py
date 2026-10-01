# -*- coding: utf-8 -*-
import os
import sys
import io
import docx

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8')

dirs_to_check = [
    "../BÁO CÁO",
    "../NHẬT KÝ",
    "../PROJECT_RAG",
    "."
]

docx_files = []
for d in dirs_to_check:
    if os.path.exists(d):
        for root, dirs, files in os.walk(d):
            for f in files:
                if f.endswith(".docx") and not f.startswith("~$"):
                    docx_files.append(os.path.join(root, f))

print(f"Checking {len(docx_files)} docx files...")

for path in docx_files:
    try:
        doc = docx.Document(path)
        found = False
        
        # Check paragraphs
        for p in doc.paragraphs:
            if "tân tạo" in p.text.lower() or "tan tao" in p.text.lower():
                found = True
                print(f"[FOUND P] in {path}: {p.text[:100]}...")
                
        # Check tables
        for t in doc.tables:
            for r in t.rows:
                for c in r.cells:
                    if "tân tạo" in c.text.lower() or "tan tao" in c.text.lower():
                        found = True
                        print(f"[FOUND TABLE] in {path}: {c.text[:100]}...")
                        
        # Check sections/headers/footers
        for s in doc.sections:
            for p in s.header.paragraphs:
                if "tân tạo" in p.text.lower() or "tan tao" in p.text.lower():
                    found = True
                    print(f"[FOUND HEADER] in {path}: {p.text[:100]}...")
            for p in s.footer.paragraphs:
                if "tân tạo" in p.text.lower() or "tan tao" in p.text.lower():
                    found = True
                    print(f"[FOUND FOOTER] in {path}: {p.text[:100]}...")
    except Exception as e:
        print(f"Error reading {path}: {e}")
