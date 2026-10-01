# -*- coding: utf-8 -*-
import os
import sys
import io
import docx

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8')

target_folder = r"g:\SMLab_KHKT\DE_TAI\RAG_SINH_HOC"

for root, dirs, files in os.walk(target_folder):
    if "venv" in root or ".git" in root:
        continue
    for f in files:
        if f.endswith(".docx") and not f.startswith("~$"):
            full_path = os.path.join(root, f)
            try:
                doc = docx.Document(full_path)
                count = 0
                for p in doc.paragraphs:
                    if "tân tạo" in p.text.lower():
                        count += 1
                for t in doc.tables:
                    for r in t.rows:
                        for c in r.cells:
                            if "tân tạo" in c.text.lower():
                                count += 1
                if count > 0:
                    print(f"File: {full_path} -> {count} occurrences", flush=True)
            except Exception as e:
                pass
