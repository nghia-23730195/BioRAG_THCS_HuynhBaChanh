import os
import docx

matches = []

for root, dirs, files in os.walk("."):
    if "venv" in root or ".git" in root or "__pycache__" in root:
        continue
    for f in files:
        path = os.path.join(root, f)
        ext = os.path.splitext(f)[1].lower()
        if ext in [".txt", ".md", ".py", ".json", ".html", ".htm"]:
            try:
                with open(path, "r", encoding="utf-8", errors="ignore") as file:
                    for i, line in enumerate(file, 1):
                        if "tân tạo" in line.lower() or "tan tao" in line.lower() or "giáo án" in line.lower():
                            matches.append((path, i, line.strip()))
            except Exception as e:
                pass
        elif ext == ".docx":
            try:
                doc = docx.Document(path)
                for i, p in enumerate(doc.paragraphs, 1):
                    if "tân tạo" in p.text.lower() or "tan tao" in p.text.lower() or "giáo án" in p.text.lower():
                        matches.append((path, f"P{i}", p.text.strip()))
                for t in doc.tables:
                    for r in t.rows:
                        for c in r.cells:
                            if "tân tạo" in c.text.lower() or "tan tao" in c.text.lower() or "giáo án" in c.text.lower():
                                matches.append((path, "Table", c.text.strip()))
            except Exception as e:
                pass

print(f"Total matches found: {len(matches)}")
for m in matches:
    print(f"{m[0]} [{m[1]}]: {m[2]}")
