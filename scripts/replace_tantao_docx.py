# -*- coding: utf-8 -*-
import os
import sys
import io
import docx

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8')

for path in ["test_40_out.docx", "../BioRAG_THCS_TanTaoA/test_40_out.docx"]:
    if os.path.exists(path):
        doc = docx.Document(path)
        modified = False
        
        for p in doc.paragraphs:
            if "TÂN TẠO" in p.text:
                p.text = p.text.replace("TRƯỜNG THCS TÂN TẠO", "TRƯỜNG THCS HUỲNH BÁ CHÁNH").replace("TÂN TẠO", "HUỲNH BÁ CHÁNH")
                modified = True
            elif "Tân Tạo" in p.text:
                p.text = p.text.replace("Tân Tạo", "Đà Nẵng")
                modified = True

        for t in doc.tables:
            for r in t.rows:
                for c in r.cells:
                    for p in c.paragraphs:
                        if "TÂN TẠO" in p.text:
                            p.text = p.text.replace("TRƯỜNG THCS TÂN TẠO", "TRƯỜNG THCS HUỲNH BÁ CHÁNH").replace("TÂN TẠO", "HUỲNH BÁ CHÁNH")
                            modified = True
                        elif "Tân Tạo" in p.text:
                            p.text = p.text.replace("Tân Tạo", "Đà Nẵng")
                            modified = True

        if modified:
            doc.save(path)
            print(f"[OK] Đã cập nhật file: {path}")
        else:
            print(f"[NO-OP] Không tìm thấy chuỗi trong {path}")
