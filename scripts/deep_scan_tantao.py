# -*- coding: utf-8 -*-
import os
import sys
import io
import docx

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8')

target_folder = r"g:\SMLab_KHKT\DE_TAI\RAG_SINH_HOC"

for root, dirs, files in os.walk(target_folder):
    if "venv" in root or ".git" in root or "__pycache__" in root:
        continue
    for f in files:
        full_path = os.path.join(root, f)
        ext = os.path.splitext(f)[1].lower()
        
        # Check text/docx/doc
        if ext in [".docx", ".doc", ".txt", ".md", ".json", ".html", ".py"]:
            if ext == ".docx":
                try:
                    doc = docx.Document(full_path)
                    all_text = []
                    for p in doc.paragraphs:
                        all_text.append(p.text)
                    for t in doc.tables:
                        for r in t.rows:
                            for c in r.cells:
                                all_text.append(c.text)
                    combined = " ".join(all_text)
                    if any(w in combined.lower() for w in ["tân tạo", "tan tao"]):
                        print(f"[DOCX MATCH] {full_path}")
                except Exception as e:
                    pass
            else:
                try:
                    with open(full_path, "r", encoding="utf-8", errors="ignore") as fl:
                        content = fl.read()
                        if any(w in content.lower() for w in ["tân tạo", "tan tao"]):
                            print(f"[TEXT MATCH] {full_path}")
                except Exception as e:
                    pass
