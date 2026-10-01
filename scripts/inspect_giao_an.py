# -*- coding: utf-8 -*-
import os
import sys
import io
import docx

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8')

def inspect_docx(filepath):
    doc = docx.Document(filepath)
    print(f"=== File: {filepath} ===")
    for i, p in enumerate(doc.paragraphs):
        if any(w in p.text.lower() for w in ["tân tạo", "huỳnh bá chánh", "giáo án", "kế hoạch bài dạy", "trường thcs"]):
            print(f"P{i+1}: {p.text}")
    for ti, t in enumerate(doc.tables):
        for ri, r in enumerate(t.rows):
            for ci, c in enumerate(r.cells):
                if any(w in c.text.lower() for w in ["tân tạo", "huỳnh bá chánh", "giáo án", "kế hoạch bài dạy", "trường thcs"]):
                    print(f"Table{ti+1} [R{ri+1}C{ci+1}]: {c.text.strip()}")

for f in ["test_40_out.docx", "../BÁO CÁO/bao_cao_final_update.docx", "../BÁO CÁO/bao_cao_Tân_Tạo_A.docx", "../BÁO CÁO/báo cáo final_HBC.docx", "../BÁO CÁO/BÁO CÁO ĐỀ TÀI.docx"]:
    if os.path.exists(f):
        inspect_docx(f)
