# -*- coding: utf-8 -*-
import os
import sys
import io
import docx

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8')

search_roots = [
    ".",
    "..",
]

found_files = []

for root_dir in search_roots:
    for root, dirs, files in os.walk(root_dir):
        if "venv" in root or ".git" in root or "__pycache__" in root or "node_modules" in root:
            continue
        for f in files:
            path = os.path.join(root, f)
            ext = os.path.splitext(f)[1].lower()
            if ext in [".docx", ".doc", ".txt", ".md", ".json", ".html", ".py"]:
                if ext == ".docx":
                    try:
                        doc = docx.Document(path)
                        full_text = " ".join([p.text for p in doc.paragraphs])
                        for t in doc.tables:
                            for r in t.rows:
                                for c in r.cells:
                                    full_text += " " + c.text
                        if "tân tạo" in full_text.lower() or "tan tao" in full_text.lower():
                            found_files.append((path, "docx"))
                    except Exception as e:
                        pass
                else:
                    try:
                        with open(path, "r", encoding="utf-8", errors="ignore") as fl:
                            content = fl.read()
                            if "tân tạo" in content.lower() or "tan tao" in content.lower():
                                found_files.append((path, ext))
                    except Exception as e:
                        pass

print(f"Total matching files found: {len(found_files)}")
for p, t in found_files:
    print(f"FOUND: {p} ({t})")
