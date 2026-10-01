import os
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
sys.path.insert(0, os.path.abspath("."))
import docx
from src.app.plan_5512_builder import build_lesson_plan_5512_doc

lesson_data = {
    "id": "k7-b2",
    "grade": 7,
    "title": "Nguyên tử",
    "number": "Bài 2",
    "duration_periods": 2,
    "source_book": "SGK KHTN 7 KNTT - Bài 2 - Trang 14–18"
}

doc = build_lesson_plan_5512_doc(lesson_data)
doc.save("scratch/test_k7_b2.docx")

doc_read = docx.Document("scratch/test_k7_b2.docx")
print("=== PARAGRAPHS ===")
for i, p in enumerate(doc_read.paragraphs[:5]):
    print(f"P{i+1}: {p.text}")

print("=== TABLES ===")
for ti, t in enumerate(doc_read.tables[:3]):
    for ri, r in enumerate(t.rows):
        for ci, c in enumerate(r.cells):
            print(f"Table{ti+1}[R{ri+1}C{ci+1}]: {c.text.strip()}")
