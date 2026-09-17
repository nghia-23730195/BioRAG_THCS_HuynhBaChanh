# -*- coding: utf-8 -*-
"""
Smart Exam Upload Parser for BioRAG (DOCX, TXT, Raw Text Paste)
Extracts questions, choices A-D, answers, and explanations from teacher-formatted exams.
"""

import io
import json
import re
import sys
import time
import zipfile
from xml.etree import ElementTree


def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extract plain text and tables from Word (.docx) file bytes.
    Uses python-docx if installed, otherwise falls back to pure zipfile + XML parser.
    """
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        lines = []
        for p in doc.paragraphs:
            txt = p.text.strip()
            if txt:
                lines.append(txt)
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)
    except Exception:
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
                xml_content = zf.read('word/document.xml')
                tree = ElementTree.fromstring(xml_content)
                paragraphs = []
                for p in tree.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'):
                    texts = [node.text for node in p.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t') if node.text]
                    if texts:
                        paragraphs.append("".join(texts).strip())
                return "\n".join(paragraphs)
        except Exception as err:
            return f"Lỗi trích xuất tệp Word .docx: {err}"


def clean_option_text(text: str) -> str:
    """Clean trailing answer tags or explanations attached to option text."""
    parts = re.split(r'\n\s*(?:Đáp án|ĐA|Chọn|Key|Giải thích|Hướng dẫn giải|Lời giải)[\s:\-_]', text, flags=re.IGNORECASE)
    cleaned = parts[0]
    cleaned = re.split(r'\s+(?:Đáp án|ĐA|Chọn|Key)[\s:\-_]+[A-D]\b', cleaned, flags=re.IGNORECASE)[0]
    return cleaned.strip()


def parse_exam_text(raw_text: str, default_grade: int = 7, default_duration: int = 15, custom_title: str = "") -> dict:
    """
    Intelligently parse Vietnamese exam text into structured quiz format.
    Supports standard question patterns, inline/bottom answer keys, and explanations.
    """
    if not raw_text or not raw_text.strip():
        return {
            "error": "Văn bản đề thi trống. Vui lòng dán đề thi hoặc chọn tệp Word/Text.",
            "questions": []
        }

    lines = [l.strip() for l in raw_text.splitlines() if l.strip()]

    # 1. Detect Metadata: Title, Grade, Duration
    title = custom_title.strip() if custom_title and custom_title.strip() else ""
    grade = int(default_grade) if default_grade else 7
    duration = int(default_duration) if default_duration else 15

    for line in lines[:10]:
        m_g = re.search(r'(?:Lớp|KHTN|Khối)\s*([6-9])', line, re.IGNORECASE)
        if m_g:
            try:
                grade = int(m_g.group(1))
            except ValueError:
                pass

        m_d = re.search(r'(\d+)\s*phút', line, re.IGNORECASE)
        if m_d:
            try:
                duration = int(m_d.group(1))
            except ValueError:
                pass

        if not title:
            lower = line.lower()
            if any(k in lower for k in ['đề kiểm tra', 'bài kiểm tra', 'đề thi', 'khảo sát', 'ôn tập', 'chủ đề:']):
                clean_t = re.sub(r'^(?:ĐỀ\s+BÀI|MÃ\s+ĐỀ|ĐỀ\s+SỐ\s*\d+)[\s:\-]*', '', line, flags=re.IGNORECASE).strip()
                if len(clean_t) > 5:
                    title = clean_t

    if not title:
        title = f"Đề kiểm tra KHTN {grade} ({duration} phút)"

    # 2. Extract Answer Key section if present at the end of the text
    answer_key_map = {}
    key_match = re.search(r'(?:(?:BẢNG|PHẦN)\s+)?(?:ĐÁP\s*ÁN|ANSWER\s*KEY).*?(?:\n|$)([\s\S]*)$', raw_text, re.IGNORECASE)
    if key_match:
        key_block = key_match.group(1)
        pairs = re.findall(r'(?:Câu\s*)?(\d+)[\.\s:\-_|]+([A-D])\b', key_block, re.IGNORECASE)
        for q_num, ans in pairs:
            answer_key_map[int(q_num)] = ord(ans.upper()) - 65

    # 3. Find Question split points
    q_pattern = r'(?:^|\n)\s*(?:Câu|Bài)\s*(\d+)[\.:\s\-]+'
    split_indices = []
    for m in re.finditer(q_pattern, raw_text, re.IGNORECASE):
        split_indices.append((m.start(), int(m.group(1)), m.group(0)))

    if not split_indices or len(split_indices) < 2:
        alt_pattern = r'(?:^|\n)\s*(\d+)[\.\)]\s+'
        alt_indices = []
        for m in re.finditer(alt_pattern, raw_text):
            alt_indices.append((m.start(), int(m.group(1)), m.group(0)))
        if len(alt_indices) >= len(split_indices):
            split_indices = alt_indices

    if not split_indices:
        return {
            "error": "Không nhận diện được cấu trúc câu hỏi (Ví dụ: 'Câu 1:', '1.'). Vui lòng kiểm tra lại định dạng đề thi.",
            "questions": []
        }

    split_indices.sort(key=lambda x: x[0])

    questions = []
    for i, (start_pos, q_num, prefix) in enumerate(split_indices):
        end_pos = split_indices[i + 1][0] if i + 1 < len(split_indices) else len(raw_text)
        if i + 1 == len(split_indices) and key_match and key_match.start() > start_pos:
            end_pos = key_match.start()

        block = raw_text[start_pos:end_pos].strip()

        block_body = block
        if block.startswith(prefix.strip()):
            block_body = block[len(prefix.strip()):].lstrip('.:- \t\n')

        opt_pattern = r'(?:^|\s+)([A-D])[\.\)]\s+'
        opt_matches = list(re.finditer(opt_pattern, block_body))

        q_text = ""
        options = ["", "", "", ""]
        ans_idx = answer_key_map.get(q_num, -1)
        expl = ""

        if len(opt_matches) >= 4:
            q_text = block_body[:opt_matches[0].start()].strip()
            for o_idx in range(4):
                o_start = opt_matches[o_idx].end()
                o_end = opt_matches[o_idx + 1].start() if o_idx < 3 else len(block_body)
                raw_opt = block_body[o_start:o_end].strip()

                if '*' in raw_opt or '(đúng)' in raw_opt.lower() or '(dung)' in raw_opt.lower():
                    ans_idx = o_idx
                    raw_opt = re.sub(r'[\*]|(?:\(đúng\))|(?:\(dung\))', '', raw_opt, flags=re.IGNORECASE).strip()

                options[o_idx] = clean_option_text(raw_opt)
        else:
            q_text = block_body
            options = ["Lựa chọn A", "Lựa chọn B", "Lựa chọn C", "Lựa chọn D"]

        m_ans = re.search(r'(?:Đáp án|ĐA|Chọn|Key)[\s:\-_]+([A-D])\b', block, re.IGNORECASE)
        if m_ans:
            ans_idx = ord(m_ans.group(1).upper()) - 65

        m_exp = re.search(r'(?:Giải thích|Hướng dẫn giải|Lời giải)[\s:\-_]+([^\n]+)', block, re.IGNORECASE)
        if m_exp:
            expl = m_exp.group(1).strip()

        if ans_idx < 0 or ans_idx > 3:
            ans_idx = 0

        diff_idx = len(questions) % 4
        diff_labels = ["Nhận biết", "Thông hiểu", "Vận dụng", "Nhận biết"]

        questions.append({
            "id": len(questions) + 1,
            "difficulty": diff_labels[diff_idx],
            "question": q_text or f"Câu hỏi {q_num}",
            "options": options,
            "answer_index": ans_idx,
            "explanation": expl or "Dựa trên kiến thức bài học trong chương trình KHTN Kết nối tri thức.",
            "source": f"KHTN {grade} (Đề tải lên)"
        })

    exam_id = f"custom-upload-{int(time.time())}"
    exam_type = "15p" if duration <= 20 else ("midterm" if duration <= 60 else "final")
    type_label = f"Kiểm tra {duration} phút" if duration <= 20 else (f"Kiểm tra Giữa kì ({duration}p)" if duration <= 60 else f"Kiểm tra Cuối kì ({duration}p)")

    return {
        "id": exam_id,
        "grade": grade,
        "type": exam_type,
        "type_label": type_label,
        "title": title,
        "topic": title,
        "duration_minutes": duration,
        "questions_count": len(questions),
        "description": f"Đề thi tải lên bởi người dùng ({len(questions)} câu hỏi trắc nghiệm).",
        "user_uploaded": True,
        "badge": "👤 Đề tải lên",
        "questions": questions
    }
